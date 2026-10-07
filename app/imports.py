import csv
import hashlib
import io
import json
import re
from collections import Counter
from datetime import date, datetime
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .models import MODELS

MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 20000
SENSITIVE = re.compile(r"password|passwd|secret|token|credential|帳密|密碼|金鑰", re.I)
ALIASES = {
    "person": ["姓名", "人員", "person"],
    "date": ["工作日期", "日期", "date"],
    "hours": ["實際時間", "工時", "hours"],
    "role": ["角色", "role"],
    "category": ["工作分類", "分類", "category"],
    "content": ["工作內容", "內容", "content"],
    "progress": ["進度說明", "progress"],
    "title": ["功能作業名稱", "標題", "名稱", "title"],
    "code": ["功能編號", "代碼", "code"],
    "system": ["系統名稱", "系統", "system"],
    "description": ["功能說明", "說明", "description"],
    "control_ref": ["對應內部控管項次", "control_ref"],
    "source_marker": ["截圖", "檢視標記", "source_marker"],
    "notes": ["問題", "備註", "notes"],
    "source_id": ["來源ID", "source_id"],
    "owner": ["負責人", "owner"],
    "amount": ["單價", "金額", "amount"],
    "purpose": ["用途", "purpose"],
    "unit": ["單位", "unit"],
    "start": ["有效起日", "start"],
    "end": ["有效迄日", "end"],
    "due": ["期限", "due"],
    "status": ["狀態", "status"],
    "kind": ["類型", "kind"],
    "priority": ["優先級", "priority"],
    "action": ["處置", "action"],
}


def scalar(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return "" if v is None else str(v)


def date_value(value):
    value = str(value).strip()
    if not value:
        return None
    parts = re.split(r"[/.-]", value[:10])
    if len(parts) == 3:
        year = int(parts[0])
        return date(year + 1911 if year < 1911 else year, int(parts[1]), int(parts[2])).isoformat()
    return value


def read_tables(filename, content):
    if len(content) > MAX_BYTES:
        raise ValueError("檔案超過8 MB上限")
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = []
        for row in reader:
            rows.append(row)
            if len(rows) > MAX_ROWS + 1:
                raise ValueError("資料超過列數上限")
        return {"CSV": rows}
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("僅接受.xlsx或UTF-8 CSV")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 5000
                or sum(i.file_size for i in archive.infolist()) > 60 * 1024 * 1024
            ):
                raise ValueError("壓縮檔展開量超過上限")
            if any("vbaProject" in i.filename or "externalLinks" in i.filename for i in archive.infolist()):
                raise ValueError("不接受巨集或外部連結")
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True, keep_links=False)
        tables = {}
        for sheet in workbook:
            sheet.reset_dimensions()
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([scalar(v) for v in row[:100]])
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError("資料超過列數上限")
            tables[sheet.title] = rows
        workbook.close()
        return tables
    except BadZipFile as exc:
        raise ValueError("無法讀取此Excel，請確認格式或權限保護") from exc


def normalize_header(value):
    return re.sub(r"\s+", "", scalar(value)).casefold()


def suggested_mapping(kind, headers):
    fields = MODELS[kind].model_fields
    result = {}
    for field, names in ALIASES.items():
        if field not in fields:
            continue
        for i, header in enumerate(headers):
            if normalize_header(header) in [normalize_header(n) for n in names] and not SENSITIVE.search(
                header
            ):
                result[field] = str(i)
                break
    return result


def sanitize_headers(headers):
    return [
        {"index": str(i), "name": scalar(h), "excluded": not scalar(h) or bool(SENSITIVE.search(scalar(h)))}
        for i, h in enumerate(headers)
    ]


def prepare(kind, project_id, rows, mapping, marker_map=None):
    if kind not in {"times", "rates", "deliverables", "issues", "works", "payments"}:
        raise ValueError("此資料類型不支援匯入")
    if kind == "rates":
        mapping = {field: column for field, column in mapping.items() if field != "tax_basis"}
    headers, records, errors = rows[0] if rows else [], [], []
    fields = MODELS[kind].model_fields
    allowed = {str(h["index"]) for h in sanitize_headers(headers) if not h["excluded"]}
    if any(f not in fields or f in {"project_id", "number"} or str(col) not in allowed
           for f, col in mapping.items()):
        raise ValueError("映射含未知欄位或敏感／未命名來源欄")
    marker_map = marker_map or {"OK": "complete", "?": "question", "N/A": "not_applicable"}
    if any(v not in {"complete", "question", "not_applicable", "unknown"} for v in marker_map.values()):
        raise ValueError("檢視標記映射錯誤")
    occurrences = Counter()
    duplicates = 0
    for number, row in enumerate(rows[1:], 2):
        if not any(scalar(v).strip() for v in row):
            continue
        payload = {f: scalar(row[int(col)]) if int(col) < len(row) else "" for f, col in mapping.items()}
        # Credentials in free-text cells are not echoed in previews or stored in records.
        payload = {f: "[敏感內容已排除]" if SENSITIVE.search(v) else v for f, v in payload.items()}
        original_fields = list(payload)
        payload["project_id"] = project_id
        try:
            for f in ["date", "start", "end", "due", "checked_on"]:
                if f in payload:
                    payload[f] = date_value(payload[f])
            for f in ["hours", "amount", "invoiced", "received"]:
                if f in payload:
                    payload[f] = payload[f].replace(",", "")
            if kind == "deliverables":
                if not payload.get("title"):
                    payload["title"] = payload.get("code", "")
                    original_fields.append("title")
                marker = marker_map.get(payload.get("source_marker", ""), "unknown")
                payload["review"] = marker if marker in {"complete", "question"} else "unknown"
                if marker == "not_applicable":
                    payload["applicable"] = False
                    payload["applicability_reason"] = "來源判定無須查核；保留原始標記"
            validated = MODELS[kind].model_validate(payload).model_dump(mode="json")
            identity_body = json.dumps({f: validated.get(f) for f in sorted(original_fields)}, sort_keys=True)
            digest = hashlib.sha256(identity_body.encode()).hexdigest()
            occurrences[digest] += 1
            if occurrences[digest] > 1:
                duplicates += 1
            source_key = payload.get("source_id") or f"{digest}:{occurrences[digest]}"
            if kind == "times":
                validated["source_id"] = source_key
            records.append(
                {"row": number, "data": validated, "source_key": source_key, "fields": original_fields}
            )
        except (ValueError, TypeError) as exc:
            # Do not echo validation input values (they can be private).
            reason = "欄位型別、必填值、日期或數值不正確"
            if hasattr(exc, "errors"):
                reason += "：" + ", ".join(str(e["loc"][0]) if e["loc"] else "日期範圍" for e in exc.errors())
            errors.append({"row": number, "reason": reason})
    return {
        "records": records,
        "errors": errors,
        "duplicate_candidates": duplicates,
        "row_count": len(records) + len(errors),
    }
