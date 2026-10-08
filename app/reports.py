import csv
import io
import math
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from .gantt_image import render_issue_gantt


def safe_text(value):
    value = "" if value is None else str(value)
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def export_table(rows, fmt="csv"):
    columns = sorted({k for row in rows for k in row if k not in {"created", "updated"}})
    if fmt == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Records"
        sheet.append(columns)
        for row in rows:
            sheet.append([safe_text(row.get(k)) for k in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(columns)
    writer.writerows([[safe_text(row.get(k)) for k in columns] for row in rows])
    return stream.getvalue().encode("utf-8-sig")


def make_snapshot(project, summary, issues, works, deliverables, external=False, title="", scenario=None):
    if external:
        # Build a new public-facing snapshot from an explicit whitelist.
        return {
            "project": {k: project[k] for k in ["name", "code", "status", "summary"]},
            "summary": {},
            "issues": [
                {
                    **{k: i[k] for k in ["title", "kind", "status", "priority"]},
                    "number": i.get("number", ""),
                    "due": i.get("due"),
                    "created": i.get("created", ""),
                    "status_history": [
                        {k: event[k] for k in ("at", "status")}
                        for event in i.get("status_history", [])
                    ],
                }
                for i in issues
            ],
            "works": [{k: i[k] for k in ["title", "kind", "due", "status"]} for i in works],
            "deliverables": [
                {k: i[k] for k in ["title", "code", "review", "applicable", "result"]} for i in deliverables
            ],
            "at": datetime.now(timezone(timedelta(hours=8))).isoformat(),
            "title": title,
            "external": True,
            "scenario": None,
        }
    return {
        "project": project,
        "summary": summary,
        "issues": issues,
        "works": works,
        "deliverables": deliverables,
        "at": datetime.now(timezone(timedelta(hours=8))).isoformat(),
        "title": title,
        "external": False,
        "scenario": scenario,
    }


def make_pptx(snapshot, sections):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    navy, teal = RGBColor(24, 38, 54), RGBColor(19, 137, 128)

    def clipped(content, limit):
        result, units = "", 0
        for char in str(content):
            units += 1 if ord(char) > 127 else 0.6
            if units > limit:
                return result + "…"
            result += char
        return result

    def slide(title):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        text(s, clipped(title, 27), 0.65, 0.55, 12, 0.7, 30, navy)
        s.notes_slide.notes_text_frame.text = str(title)
        text(
            s,
            f"{clipped(snapshot['project']['name'], 65)}  /  {snapshot['at'][:10]}",
            0.65,
            7.05,
            12,
            0.25,
            10,
            teal,
        )
        return s

    def text(s, content, x, y, w, h, size=20, color=navy):
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.text_frame.word_wrap = True
        p = box.text_frame.paragraphs[0]
        p.text = str(content)
        p.font.size, p.font.name, p.font.color.rgb = Pt(size), "Microsoft JhengHei", color
        return box

    def issue_timeline(rows):
        if not rows:
            return
        dated = [row for row in rows if row.get("created") and row.get("due")]
        starts = [row["created"][:10] for row in dated]
        ends = [max(row["due"], row["created"][:10]) for row in dated]
        if not starts:
            s = slide("事項追蹤甘特圖")
            text(s, "事項尚未設定期限，無法繪製追蹤區間", 0.8, 1.8, 11.5, 0.6, 18)
            return
        low, high = min(starts), max(ends)
        pages = [rows[i:i + 8] for i in range(0, len(rows), 8)]
        for page_index, page_rows in enumerate(pages):
            s = slide("事項追蹤甘特圖" + (f" ({page_index + 1}/{len(pages)})" if len(pages) > 1 else ""))
            picture = io.BytesIO(render_issue_gantt(page_rows, low, high, snapshot["external"]))
            s.shapes.add_picture(picture, Inches(0.7), Inches(1.3), width=Inches(11.9), height=Inches(5.4))

    def table(title, columns, rows):
        if columns[0][1] == "number":
            widths = [1.7, 3.6, 2.4, 1.4, 1.4, 1.4] if any(
                key == "owner" for _, key in columns
            ) else [1.7, 4.9, 1.8, 1.8, 1.7]
        elif any(key == "owner" for _, key in columns):
            widths = [4.1] + [2.4 if key == "owner" else 1.8 for _, key in columns[1:]]
        else:
            widths = [5.5] + [(11.9 - 5.5) / (len(columns) - 1)] * (len(columns) - 1)
        labels = {
            "open": "待處理",
            "in_progress": "處理中",
            "resolved": "已解決",
            "closed": "已結案",
            "issue": "議題",
            "risk": "風險",
            "change": "變更",
            "decision": "決策",
            "task": "待辦",
            "milestone": "里程碑",
            "todo": "待辦",
            "doing": "進行中",
            "done": "完成",
            "unknown": "未確認",
            "complete": "檢視完成",
            "question": "有疑問",
            "pass": "通過",
            "fail": "失敗",
            "not_run": "未執行",
            "blocked": "阻塞",
            "low": "低",
            "medium": "中",
            "high": "高",
            "critical": "緊急",
        }

        def value(item, key):
            v = item.get(key)
            if key == "owner" and item.get("owner_display"):
                v = item["owner_display"]
            if key == "owner" and not v:
                return "未指定"
            if key == "owner" and not item.get("owner_display") and item.get("owner_alias"):
                v = item["owner_alias"]
            return (
                "待確認"
                if v is None
                else "是"
                if v is True
                else "否"
                if v is False
                else labels.get(str(v), str(v))
            )

        pages, heights, page, page_heights, used = [], [], [], [], 0.5
        for item in rows:
            lines = max(
                sum(
                    max(
                        1,
                        math.ceil(
                            sum(1 if ord(c) > 127 else 0.6 for c in line) / max(1, (widths[i] - 0.2) / 0.21)
                        ),
                    )
                    for line in value(item, key).split("\n")
                )
                for i, (_, key) in enumerate(columns)
            )
            height = max(0.5, lines * 0.25 + 0.12)
            if page and used + height > 4.8:
                pages.append(page)
                heights.append(page_heights)
                page, page_heights, used = [], [], 0.5
            page.append(item)
            page_heights.append(height)
            used += height
        pages.append(page)
        heights.append(page_heights)
        for idx, page in enumerate(pages):
            s = slide(title + (f" ({idx + 1}/{len(pages)})" if len(pages) > 1 else ""))
            if not page:
                text(s, "此報告期間沒有相關紀錄", 0.7, 1.7, 11.8, 1)
                continue
            t = s.shapes.add_table(
                len(page) + 1,
                len(columns),
                Inches(0.7),
                Inches(1.7),
                Inches(11.9),
                Inches(0.5 + sum(heights[idx])),
            ).table
            t.rows[0].height = Inches(0.5)
            for row, height in enumerate(heights[idx], 1):
                t.rows[row].height = Inches(height)
            for col, (label, key) in enumerate(columns):
                t.columns[col].width = Inches(widths[col])
                t.cell(0, col).text = label
                for row, item in enumerate(page, 1):
                    t.cell(row, col).text = value(item, key)
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        p.font.size, p.font.name = Pt(15), "Microsoft JhengHei"

    s = slide(snapshot.get("title") or "專案進度報告")
    text(s, clipped(snapshot["project"]["name"], 21), 0.7, 2, 11.7, 1, 38, teal)
    s.notes_slide.notes_text_frame.text += "\n" + snapshot["project"]["name"]
    summary_text = snapshot["project"].get("summary", "") or "尚未填寫專案摘要"
    text(s, clipped(summary_text, 150), 0.7, 3.2, 11.7, 2, 22)
    if not snapshot["external"]:
        owner = snapshot["project"].get("owner_alias") or snapshot["project"].get("owner") or "未設定"
        text(s, f"專案負責人：{clipped(owner, 70)}", 0.7, 5.65, 11.7, 0.4, 18)
    text(s, "對外版" if snapshot["external"] else "內部管理版", 0.7, 6.2, 11, 0.5, 16)
    if len(summary_text) > 150:
        for begin in range(0, len(summary_text), 450):
            s = slide("專案摘要")
            text(s, summary_text[begin : begin + 450], 0.7, 1.7, 11.7, 4.8, 20)
    if "effort" in sections and not snapshot["external"]:
        summary = snapshot["summary"]
        s = slide("資源投入與成本")
        text(
            s,
            f"累計工時  {summary['hours']} 小時\n預計完成成本  {summary['eac'] or '待估'}\n"
            f"全期已投入成本  {summary['actual_cost'] or '待估'}\n剩餘成本  {summary.get('etc') or '待估'}\n"
            f"預估餘額  {summary['profit'] or '待估'}\n"
            f"缺少成本單價  {summary['missing_rate_rows']} 筆",
            0.7,
            1.6,
            4,
            3.7,
            20,
        )
        if summary["monthly"]:
            data = CategoryChartData()
            data.categories = [r["name"] for r in summary["monthly"]]
            data.add_series("工時", [r["hours"] for r in summary["monthly"]])
            s.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.7), Inches(7.4), Inches(4.6), data
            )
    if "issues" in sections:
        issue_timeline(snapshot["issues"])
        table(
            "議題與決策",
            [("事項編號", "number"), ("事項", "title"), ("類型", "kind"), ("優先級", "priority"), ("狀態", "status")]
            if snapshot["external"] else
            [("事項編號", "number"), ("事項", "title"), ("負責人", "owner"), ("類型", "kind"), ("優先級", "priority"), ("狀態", "status")],
            snapshot["issues"],
        )
    if "works" in sections:
        table(
            "待辦與里程碑",
            [("工作", "title"), ("類型", "kind"), ("期限", "due"), ("狀態", "status")]
            if snapshot["external"] else
            [("工作", "title"), ("負責人", "owner"), ("類型", "kind"), ("期限", "due"), ("狀態", "status")],
            snapshot["works"],
        )
    if "deliverables" in sections:
        table(
            "交付與查核",
            [("交付", "title"), ("初步檢視", "review"), ("需查核", "applicable"), ("正式結果", "result")]
            if snapshot["external"] else
            [("交付", "title"), ("負責人", "owner"), ("初步檢視", "review"), ("需查核", "applicable"), ("正式結果", "result")],
            snapshot["deliverables"],
        )
    if "scenario" in sections and snapshot.get("scenario") and not snapshot["external"]:
        v = snapshot["scenario"]
        s = slide("變更情境")
        text(
            s,
            f"{v['name']} ({v['status']})\n修訂收入：{v['revised_revenue'] or '待估'}\n"
            f"淨收入減少：{v['net_revenue_decrease'] or '待估'}\n"
            f"修訂完工成本：{v['revised_eac'] or '待估'}",
            0.7,
            1.7,
            11.7,
            4.5,
            26,
        )
    stream = io.BytesIO()
    prs.save(stream)
    return stream.getvalue()
