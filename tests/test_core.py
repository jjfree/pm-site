import io
import json
import sqlite3
from zipfile import ZipFile

from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches
from app.analytics import evaluate_scenario, summarize
from app.db import Store
from app.models import Project, Rate, Scenario, TimeEntry
from app.reports import export_table, make_pptx, make_snapshot
from conftest import body


def model(cls, **kw):
    return cls(**kw).model_dump(mode="json")


def test_synthetic_demo_is_complete_and_only_for_empty_database(client):
    assert client.post("/api/demo").status_code == 200
    assert len(client.get("/api/records/projects").json()) == 3
    assert len(client.get("/api/records/issues").json()) == 6
    assert len(client.get("/api/records/works").json()) == 6
    assert len(client.get("/api/records/times").json()) == 36
    assert client.post("/api/demo").status_code == 409


def test_local_security(client):
    assert client.get("/api/overview", headers={"Host": "evil.example"}).status_code == 403
    assert client.get("/api/overview", headers={"sec-fetch-site": "cross-site"}).status_code == 403
    assert (
        client.post("/api/records/projects", json={"name": "X"}, headers={"x-csrf-token": "bad"}).status_code
        == 403
    )
    assert (
        client.post(
            "/api/records/projects", json={"name": "X"}, headers={"origin": "https://evil.example"}
        ).status_code
        == 403
    )
    client.cookies.clear()
    assert client.get("/api/overview").status_code == 401


def test_crud_version_and_relations(client, project):
    pid = project["id"]
    updated = client.put(f"/api/records/projects/{pid}", json={**body(project), "summary": "Changed"})
    assert updated.json()["version"] == 2
    assert client.put(f"/api/records/projects/{pid}", json=body(project)).status_code == 409
    assert client.post("/api/records/issues", json={"project_id": "missing", "title": "X"}).status_code == 422
    issue = client.post("/api/records/issues", json={"project_id": pid, "title": "X"}).json()
    assert client.delete(f"/api/records/projects/{pid}?version=2").status_code == 409
    assert client.delete(f"/api/records/issues/{issue['id']}?version=1").status_code == 200
    assert client.delete(f"/api/records/projects/{pid}?version=2").status_code == 200


def test_project_code_required_and_issue_numbers_are_stable(client, project):
    pid = project["id"]
    assert client.post("/api/records/projects", json={
        "name": "Missing code", "tax_basis": "exclusive",
    }).status_code == 422
    assert client.post("/api/records/projects", json={
        "name": "Duplicate code", "code": "syn", "tax_basis": "exclusive",
    }).status_code == 409
    first = client.post("/api/records/issues", json={
        "project_id": pid, "title": "First", "number": "CUSTOM-9999",
    }).json()
    second = client.post("/api/records/issues", json={
        "project_id": pid, "title": "Second",
    }).json()
    assert (first["number"], second["number"]) == ("SYN-0001", "SYN-0002")
    changed = client.put(f"/api/records/issues/{second['id']}", json={
        **body(second), "number": "CUSTOM-0001", "status": "in_progress",
    }).json()
    assert changed["number"] == "SYN-0002"
    assert client.delete(f"/api/records/issues/{first['id']}?version={first['version']}").status_code == 200
    third = client.post("/api/records/issues", json={
        "project_id": pid, "title": "Third",
    }).json()
    assert third["number"] == "SYN-0003"
    assert client.put(f"/api/records/projects/{pid}", json={
        **body(project), "code": "RENAMED",
    }).status_code == 422
    report = client.post("/api/reports", json={"project_id": pid, "sections": ["issues"]}).json()
    assert [row["number"] for row in report["snapshot"]["issues"]] == ["SYN-0002", "SYN-0003"]
    pptx = Presentation(io.BytesIO(client.get(f"/api/reports/{report['id']}/pptx").content))
    tables = [shape.table for slide in pptx.slides for shape in slide.shapes if shape.has_table]
    assert tables[0].cell(0, 0).text == "事項編號"
    assert [tables[0].cell(row, 0).text for row in (1, 2)] == ["SYN-0002", "SYN-0003"]
    public = client.post("/api/reports", json={
        "project_id": pid, "external": True, "sections": ["issues"],
    }).json()
    assert [row["number"] for row in public["snapshot"]["issues"]] == ["SYN-0002", "SYN-0003"]


def test_legacy_project_issues_receive_numbers_when_code_is_set(client):
    store = client.app.state.store
    with store.connect() as db:
        legacy = store.write("projects", {"name": "Legacy", "code": "", "tax_basis": "exclusive"}, db=db)
        for index, title in enumerate(("Old first", "Old second"), 1):
            db.execute("INSERT INTO records VALUES(?,?,?,?,?,?,?)", (
                f"legacy-{index}", "issues", legacy["id"],
                json.dumps({"project_id": legacy["id"], "title": title, "status": "open"}),
                1, f"2025-01-0{index}T00:00:00+00:00", f"2025-01-0{index}T00:00:00+00:00",
            ))
    assert client.post("/api/records/issues", json={
        "project_id": legacy["id"], "title": "New",
    }).status_code == 422
    updated = client.put(f"/api/records/projects/{legacy['id']}", json={
        **body(legacy), "code": "OLD",
    })
    assert updated.status_code == 200
    issues = client.get(f"/api/records/issues?project_id={legacy['id']}").json()
    assert [issue["number"] for issue in issues] == ["OLD-0001", "OLD-0002"]
    assert client.post("/api/records/issues", json={
        "project_id": legacy["id"], "title": "New",
    }).json()["number"] == "OLD-0003"


def test_issue_number_migration_backfills_and_backs_up(tmp_path):
    path = tmp_path / "legacy"
    store = Store(path)
    project = store.write("projects", {"name": "Legacy", "code": "OLD", "tax_basis": "exclusive"})
    with store.connect() as db:
        db.execute("INSERT INTO records VALUES(?,?,?,?,?,?,?)", (
            "legacy-issue", "issues", project["id"],
            json.dumps({"project_id": project["id"], "title": "Old issue", "status": "open"}),
            1, "2025-01-01T00:00:00+00:00", "2025-01-01T00:00:00+00:00",
        ))
        db.execute("PRAGMA user_version=2")
    migrated = Store(path)
    assert migrated.get("issues", "legacy-issue")["number"] == "OLD-0001"
    assert len(list((path / "backups").glob("backup-before-issue-numbers-*.sqlite3"))) == 1
    assert Store(path).get("issues", "legacy-issue")["number"] == "OLD-0001"


def test_issue_timeline_tracks_status_and_prioritizes_assignee_alias(client, project):
    pid = project["id"]
    member = client.post("/api/records/members", json={
        "project_id": pid, "person": "林小明", "alias": "Alex", "role": "PM",
    }).json()
    issue = client.post("/api/records/issues", json={
        "project_id": pid, "title": "狀態歷程事項", "owner_member_id": member["id"],
        "due": "2026-12-31", "status": "open",
    }).json()
    updated = client.put(f"/api/records/issues/{issue['id']}", json={
        **body(issue), "status": "in_progress",
    })
    assert updated.status_code == 200
    renamed = client.put(f"/api/records/members/{member['id']}", json={
        **body(member), "alias": "Alex-New",
    })
    assert renamed.status_code == 200
    timeline_row = next(row for row in client.get(
        f"/api/records/issues?project_id={pid}&aliases=true"
    ).json() if row["id"] == issue["id"])
    assert [event["status"] for event in timeline_row["status_history"]] == ["open", "in_progress"]
    assert all(event["owner_alias"] == "Alex-New" for event in timeline_row["status_history"])

    report = client.post("/api/reports", json={"project_id": pid, "sections": ["issues"]}).json()
    saved_issue = next(row for row in report["snapshot"]["issues"] if row["id"] == issue["id"])
    assert saved_issue["status_history"] == timeline_row["status_history"]
    pptx = Presentation(io.BytesIO(client.get(f"/api/reports/{report['id']}/pptx").content))
    texts = [shape.text for slide in pptx.slides for shape in slide.shapes if shape.has_text_frame]
    assert "事項追蹤甘特圖" in texts
    assert texts.count("Alex-New") == 1
    assert "Alex" not in texts
    assert "林小明" not in texts
    public = client.post("/api/reports", json={
        "project_id": pid, "external": True, "sections": ["issues"],
    }).json()
    serialized = json.dumps(public["snapshot"], ensure_ascii=False)
    assert "Alex-New" not in serialized and "林小明" not in serialized
    assert [event["status"] for event in public["snapshot"]["issues"][0]["status_history"]] == ["open", "in_progress"]


def test_gantt_shows_same_owner_once_when_free_text_is_linked(client, project):
    member = client.post("/api/records/members", json={
        "project_id": project["id"], "person": "林小明", "alias": "Ted",
    }).json()
    issue = client.post("/api/records/issues", json={
        "project_id": project["id"], "title": "Owner link", "owner": "Ted", "due": "2025-01-20",
    }).json()
    linked = client.put(f"/api/records/issues/{issue['id']}", json={
        **body(issue), "owner_member_id": member["id"],
    })
    assert linked.status_code == 200
    with client.app.state.store.connect() as db:
        db.execute("UPDATE records SET created=? WHERE id=?", ("2025-01-01T00:00:00+00:00", issue["id"]))
        audits = db.execute("SELECT id FROM audits WHERE record_id=? ORDER BY id", (issue["id"],)).fetchall()
        for audit, at in zip(audits, ("2025-01-01T00:00:00+00:00", "2025-01-05T00:00:00+00:00")):
            db.execute("UPDATE audits SET at=? WHERE id=?", (at, audit["id"]))
    history = client.get(f"/api/records/issues?project_id={project['id']}&aliases=true").json()[0]["status_history"]
    assert [event["owner_alias"] or event["owner"] for event in history] == ["Ted", "Ted"]
    report = client.post("/api/reports", json={"project_id": project["id"], "sections": ["issues"]}).json()
    pptx = Presentation(io.BytesIO(client.get(f"/api/reports/{report['id']}/pptx").content))
    texts = [shape.text for slide in pptx.slides for shape in slide.shapes if shape.has_text_frame]
    assert texts.count("Ted") == 1
    owner_badge = next(shape for slide in pptx.slides for shape in slide.shapes
                       if shape.has_text_frame and shape.text == "Ted")
    assert abs(owner_badge.left - Inches(4.35 + 4 / 20 * 8.15 + 0.02)) < Inches(0.05)


def test_validation_and_rate_overlap(client, project):
    rate = {
        "project_id": project["id"],
        "role": "Engineer",
        "amount": "3200",
        "tax_basis": "exclusive",
        "start": "2024-01-01",
        "end": "2024-12-31",
    }
    saved = client.post("/api/records/rates", json=rate)
    assert saved.status_code == 200 and "tax_basis" not in saved.json()
    assert client.post("/api/records/rates", json={**rate, "start": "2024-12-31"}).status_code == 409
    assert (
        client.post("/api/records/rates", json={**rate, "start": "2025-01-01", "end": None}).status_code
        == 200
    )
    assert client.post("/api/records/rates", json={**rate, "amount": "-1"}).status_code == 422
    assert client.post("/api/records/projects", json={"name": "Missing basis", "code": "VALIDATION"}).status_code == 422
    assert client.post("/api/records/projects", json={"name": "Unknown basis", "code": "VALIDATION", "tax_basis": "unknown"}).status_code == 422
    assert (
        client.post(
            "/api/records/projects", json={"name": "X", "code": "VALIDATION", "tax_basis": "exclusive", "start": "2025-02-01", "end": "2025-01-01"}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/records/issues",
            json={"project_id": project["id"], "title": "X", "external_url": "javascript:alert(1)"},
        ).status_code
        == 422
    )


def test_legacy_rate_tax_is_hidden_from_records_and_exports(client, project):
    store = client.app.state.store
    rate = store.write("rates", {
        "project_id": project["id"], "role": "Engineer", "amount": "3200",
        "purpose": "cost", "person": "", "unit": "day", "start": "2000-01-01",
        "end": None, "tax_basis": "exclusive",
    })
    visible = client.get(f"/api/records/rates?project_id={project['id']}").json()[0]
    assert visible["id"] == rate["id"] and "tax_basis" not in visible
    exported = client.get(f"/api/exports/rates?project_id={project['id']}").text
    assert "tax_basis" not in exported.splitlines()[0]


def test_time_role_edit_preserves_other_entries_and_cost(client, project):
    pid = project["id"]
    client.post(
        "/api/records/rates",
        json={"project_id": pid, "role": "Engineer", "amount": "4000", "tax_basis": "exclusive"},
    )
    entries = [
        client.post(
            "/api/records/times",
            json={"project_id": pid, "person": person, "date": "2025-01-01", "hours": 8, "role": role},
        ).json()
        for person, role in [("Member-A", ""), ("Member-A", "Old"), ("Member-B", "")]
    ]
    other_project = client.post("/api/records/projects", json={"name": "Other project", "code": "OTHER", "tax_basis": "exclusive"}).json()
    outside = client.post(
        "/api/records/times",
        json={"project_id": other_project["id"], "person": "Member-A", "date": "2025-01-01", "hours": 8},
    ).json()

    response = client.put(
        f"/api/records/times/{entries[0]['id']}",
        json={**body(entries[0]), "role": "Engineer"},
    )
    assert response.status_code == 200
    rows = client.get(f"/api/records/times?project_id={pid}").json()
    assert [row["role"] for row in rows] == ["Engineer", "Old", ""]
    assert rows[1]["version"] == entries[1]["version"]
    assert client.get(f"/api/records/times?project_id={other_project['id']}").json()[0]["role"] == ""
    assert client.get(f"/api/analytics/{pid}").json()["summary"]["known_labor_cost"] == "4000.00"

    rejected = client.put(
        f"/api/records/times/{entries[1]['id']}",
        json={**body(entries[1]), "role": "Engineer", "apply_role_to_person": True},
    )
    assert rejected.status_code == 422

    stale = client.put(
        f"/api/records/times/{entries[0]['id']}",
        json={**body(entries[0]), "role": "Engineer"},
    )
    assert stale.status_code == 409
    assert client.get(f"/api/records/times?project_id={pid}").json() == rows
    assert client.get(f"/api/records/times?project_id={other_project['id']}").json()[0] == outside


def test_project_members_and_roles_are_unique_and_do_not_rewrite_history(client, project):
    pid = project["id"]
    role = client.post("/api/records/roles", json={"project_id": pid, "name": "Engineer"}).json()
    assert role["active"] is True
    assert client.post("/api/records/roles", json={"project_id": pid, "name": "engineer"}).status_code == 409
    member = client.post("/api/records/members", json={"project_id": pid, "person": "Member-A", "role": "Engineer"}).json()
    assert client.post("/api/records/members", json={"project_id": pid, "person": "member-a"}).status_code == 409
    assert client.put(f"/api/records/roles/{role['id']}", json={**body(role), "name": "Renamed"}).status_code == 422
    assert client.put(f"/api/records/members/{member['id']}", json={**body(member), "person": "Renamed"}).status_code == 422
    time = client.post("/api/records/times", json={"project_id": pid, "person": "Member-A", "role": "Engineer", "date": "2025-01-01", "hours": 8}).json()
    updated = client.put(f"/api/records/members/{member['id']}", json={**body(member), "role": "PM"})
    assert updated.status_code == 200
    assert client.get(f"/api/records/times?project_id={pid}").json()[0] == time
    assert client.delete(f"/api/records/members/{member['id']}?version=2").status_code == 409
    assert client.delete(f"/api/records/roles/{role['id']}?version=1").status_code == 409
    assert client.put(f"/api/records/roles/{role['id']}", json={**body(role), "active": False}).json()["active"] is False
    other = client.post("/api/records/projects", json={"name": "Other", "code": "OTHER", "tax_basis": "exclusive"}).json()
    assert client.post("/api/records/roles", json={"project_id": other["id"], "name": "Engineer"}).status_code == 200
    assert client.post("/api/records/members", json={"project_id": other["id"], "person": "Member-A"}).status_code == 200


def test_owner_assignments_link_to_project_members_and_preserve_legacy_values(client, project):
    pid = project["id"]
    member = client.post("/api/records/members", json={
        "project_id": pid, "person": "Lin", "alias": "Alex", "role": "PM",
    }).json()
    assert client.post("/api/records/members", json={
        "project_id": pid, "person": "Other", "alias": "alex",
    }).status_code == 409
    assert client.post("/api/records/members", json={
        "project_id": pid, "person": "Alex",
    }).status_code == 409

    legacy = client.post("/api/records/issues", json={
        "project_id": pid, "title": "Legacy", "owner": "Old name",
    }).json()
    assert legacy["owner"] == "Old name" and legacy["owner_member_id"] == ""
    assigned = client.put(f"/api/records/issues/{legacy['id']}", json={
        **body(legacy), "owner_member_id": member["id"], "owner": "Spoofed",
    }).json()
    assert assigned["owner"] == "Lin" and assigned["owner_member_id"] == member["id"]
    for kind, title in (("works", "Task"), ("deliverables", "Deliverable")):
        created = client.post(f"/api/records/{kind}", json={
            "project_id": pid, "title": title, "owner_member_id": member["id"],
        }).json()
        assert created["owner"] == "Lin"
    updated_project = client.put(f"/api/records/projects/{pid}", json={
        **body(project), "owner_member_id": member["id"],
    }).json()
    assert updated_project["owner"] == "Lin"

    other = client.post("/api/records/projects", json={"name": "Other", "code": "OTHER", "tax_basis": "exclusive"}).json()
    assert client.post("/api/records/issues", json={
        "project_id": other["id"], "title": "Cross project", "owner_member_id": member["id"],
    }).status_code == 422
    inactive = client.put(f"/api/records/members/{member['id']}", json={
        **body(member), "active": False,
    }).json()
    assert client.put(f"/api/records/issues/{legacy['id']}", json={
        **body(assigned), "action": "Still assigned",
    }).status_code == 200
    assert client.post("/api/records/issues", json={
        "project_id": pid, "title": "New assignment", "owner_member_id": member["id"],
    }).status_code == 422
    assert client.delete(f"/api/records/members/{member['id']}?version={inactive['version']}").status_code == 409


def test_aliases_appear_in_records_exports_and_internal_report_only(client, project):
    pid = project["id"]
    member = client.post("/api/records/members", json={
        "project_id": pid, "person": "Lin", "alias": "Alex", "role": "PM",
    }).json()
    other = client.post("/api/records/projects", json={"name": "Other", "code": "OTHER", "tax_basis": "exclusive"}).json()
    client.post("/api/records/members", json={
        "project_id": other["id"], "person": "Lin", "alias": "Different",
    })
    client.put(f"/api/records/projects/{pid}", json={
        **body(project), "owner_member_id": member["id"],
    })
    for kind, title in (("issues", "Issue"), ("works", "Work"), ("deliverables", "Delivery")):
        client.post(f"/api/records/{kind}", json={
            "project_id": pid, "title": title, "owner_member_id": member["id"],
        })
    client.post("/api/records/issues", json={
        "project_id": pid, "title": "Legacy", "owner": "Lin",
    })
    client.post("/api/records/times", json={
        "project_id": pid, "person": "Lin", "date": "2025-01-01", "hours": 8,
    })
    client.post("/api/records/rates", json={
        "project_id": pid, "role": "PM", "person": "Lin", "amount": "100",
        "start": "2025-01-01",
    })
    assert client.get("/api/overview").json()[0]["owner_alias"] == "Alex"
    assert client.get("/api/records/projects?aliases=true").json()[0]["owner_alias"] == "Alex"
    for kind in ("issues", "works", "deliverables"):
        rows = client.get(f"/api/records/{kind}?project_id={pid}&aliases=true").json()
        assert rows[0]["owner_alias"] == "Alex"
        sheet = load_workbook(io.BytesIO(client.get(
            f"/api/exports/{kind}?project_id={pid}&fmt=xlsx"
        ).content)).active
        assert "owner_alias" in [cell.value for cell in sheet[1]]
        assert any(cell.value == "Alex" for cell in sheet[2])
        if kind == "issues":
            assert rows[1]["owner_alias"] == ""  # Free-text owners are not silently linked.
    for kind in ("times", "rates"):
        rows = client.get(f"/api/records/{kind}?project_id={pid}&aliases=true").json()
        assert rows[0]["person_alias"] == "Alex"
        sheet = load_workbook(io.BytesIO(client.get(
            f"/api/exports/{kind}?project_id={pid}&fmt=xlsx"
        ).content)).active
        assert "person_alias" in [cell.value for cell in sheet[1]]
        assert any(cell.value == "Alex" for cell in sheet[2])
    report = client.post("/api/reports", json={"project_id": pid}).json()
    assert report["snapshot"]["project"]["owner_alias"] == "Alex"
    assert report["snapshot"]["issues"][0]["owner_alias"] == "Alex"
    pptx = Presentation(io.BytesIO(client.get(f"/api/reports/{report['id']}/pptx").content))
    cells = [cell.text for slide in pptx.slides for shape in slide.shapes
             if shape.has_table for row in shape.table.rows for cell in row.cells]
    assert "Alex" in cells
    assert "Lin（Alex）" not in cells
    public = client.post("/api/reports", json={"project_id": pid, "external": True}).json()
    assert "Alex" not in json.dumps(public["snapshot"])
    public_pptx = client.get(f"/api/reports/{public['id']}/pptx").content
    with ZipFile(io.BytesIO(public_pptx)) as archive:
        assert all(b"Alex" not in archive.read(name) for name in archive.namelist() if name.endswith(".xml"))


def test_explicit_bulk_role_change_checks_preview_and_updates_atomically(client, project):
    pid = project["id"]
    client.post("/api/records/members", json={"project_id": pid, "person": "Member-A", "role": "Engineer"})
    entries = [client.post("/api/records/times", json={
        "project_id": pid, "person": "Member-A", "role": "", "date": "2025-01-01", "hours": 8,
    }).json() for _ in range(2)]
    payload = {"project_id": pid, "person": "Member-A", "role": "Engineer",
               "entries": [{"id": row["id"], "version": row["version"]} for row in entries]}
    assert client.post("/api/times/bulk-role", json={**payload, "entries": payload["entries"][:1]}).status_code == 409
    assert client.get(f"/api/records/times?project_id={pid}").json() == entries
    result = client.post("/api/times/bulk-role", json=payload)
    assert result.status_code == 200 and result.json()["updated"] == 2
    updated = client.get(f"/api/records/times?project_id={pid}").json()
    assert all(row["role"] == "Engineer" and row["version"] == 2 for row in updated)
    assert client.post("/api/times/bulk-role", json=payload).status_code == 409
    assert client.get(f"/api/records/times?project_id={pid}").json() == updated


def fixtures():
    p = model(
        Project,
        name="Synthetic",
        code="SYN",
        hours_per_day=8,
        revenue=200000,
        eac=27200,
        budget=90000,
        tax_basis="exclusive",
    )
    rates = [
        model(Rate, project_id="P", role="Engineer", amount=3200, end="2024-12-31"),
        model(Rate, project_id="P", role="Engineer", amount=4000, start="2025-01-01"),
    ]
    times = [
        model(TimeEntry, project_id="P", person="Member-A", role="Engineer", date=d, hours=8)
        for d in ["2024-12-31", "2025-01-01"]
    ]
    return p, times, rates


def test_cross_year_cost_and_filtered_scope():
    p, times, rates = fixtures()
    s = summarize(p, times, rates)
    assert s["hours"] == "16" and s["md"] == "2.00"
    assert s["actual_cost"] == "7200.00" and s["eac"] == "27200.00"
    assert s["profit"] == "172800.00"
    filtered = summarize(p, times, rates, "2025-01-01")
    assert filtered["known_labor_cost"] == "4000.00"
    assert filtered["eac"] == "27200.00" and filtered["profit"] == "172800.00"
    assert filtered["actual_cost"] == "7200.00" and filtered["budget"] == "90000"
    assert filtered["etc"] == "20000.00"
    p["tax_basis"] = "inclusive"
    assert summarize(p, times, rates)["actual_cost"] == "7200.00"


def test_missing_is_not_zero_and_legacy_rate_basis():
    p, times, rates = fixtures()
    s = summarize(p, times, rates[:1])
    assert s["known_labor_cost"] == "3200.00" and s["missing_rate_rows"] == 1
    assert s["actual_cost"] is None
    assert s["eac"] == "27200.00" and s["profit"] == "172800.00"
    rates[1]["tax_basis"] = "inclusive"
    mismatched = summarize(p, times, rates)
    assert mismatched["actual_cost"] is None and mismatched["legacy_basis_conflicts"] == 1
    rates[1]["tax_basis"] = "unknown"
    assert summarize(p, times, rates)["actual_cost"] == "7200.00"
    p["hours_per_day"] = None
    assert summarize(p, times, rates)["md"] is None
    rates[0]["unit"] = "hour"
    assert summarize(p, times[:1], rates)["known_labor_cost"] == "25600.00"


def test_person_rate_and_scenario_unknown():
    p, times, rates = fixtures()
    rates.append({**rates[1], "person": "Member-A", "amount": "4500"})
    s = summarize(p, times, rates)
    assert s["known_labor_cost"] == "7700.00"
    scenario = model(Scenario, project_id="P", name="Option", removed_value=30000)
    assert evaluate_scenario(p, scenario, rates, s)["net_revenue_decrease"] is None
    rates.append(model(Rate, project_id="P", role="Engineer", purpose="sale", amount=8000))
    scenario.update(billable_md="2", sale_role="Engineer", rate_date="2025-01-01", cost_change="1000")
    result = evaluate_scenario(p, scenario, rates, s)
    assert result["replacement_revenue"] == "16000.00" and result["net_revenue_decrease"] == "14000.00"
    assert result["revised_eac"] == "28200.00"
    scenario["additional_revenue"] = "0"
    assert evaluate_scenario(p, scenario, rates, s)["replacement_revenue"] == "0.00"


def test_backup_integrity(client, project):
    r = client.post("/api/backups").json()
    binary = client.get("/api/backups/" + r["name"])
    assert binary.status_code == 200 and len(r["sha256"]) == 64
    store = client.app.state.store
    with sqlite3.connect(store.directory / "backups" / r["name"]) as db:
        assert db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert db.execute("SELECT count(*) FROM records").fetchone()[0] == 1


def test_export_formula_injection():
    rows = [{"title": '=HYPERLINK("https://invalid.example")', "content": "  @bad", "hours": "8"}]
    assert "'=HYPERLINK" in export_table(rows).decode("utf-8-sig")
    wb = load_workbook(io.BytesIO(export_table(rows, "xlsx")))
    assert all(c.data_type != "f" for row in wb.active for c in row)


def test_report_whitelist_and_native_chart(client, project):
    client.post(
        "/api/records/issues",
        json={
            "project_id": project["id"],
            "title": "Visible",
            "owner": "Private-member",
            "description": "Private-note",
        },
    )
    client.post(
        "/api/records/times",
        json={"project_id": project["id"], "person": "Private-member", "date": "2025-01-01", "hours": 8},
    )
    response = client.post("/api/reports", json={"project_id": project["id"], "external": True}).json()
    serialized = json.dumps(response["snapshot"])
    assert (
        "Private-member" not in serialized
        and "Private-note" not in serialized
        and "revenue" not in serialized
    )
    content = client.get(f"/api/reports/{response['id']}/pptx").content
    prs = Presentation(io.BytesIO(content))
    assert len(prs.slides) >= 2
    with ZipFile(io.BytesIO(content)) as z:
        assert all(b"Private-member" not in z.read(n) for n in z.namelist() if n.endswith(".xml"))
    internal = client.post("/api/reports", json={"project_id": project["id"]}).json()
    binary = client.get(f"/api/reports/{internal['id']}/pptx").content
    with ZipFile(io.BytesIO(binary)) as z:
        assert any(n.startswith("ppt/charts/chart") for n in z.namelist())
    client.put(f"/api/records/projects/{project['id']}", json={**body(project), "name": "Renamed"})
    assert client.get(f"/api/reports/{internal['id']}/pptx").content != b""
    with client.app.state.store.connect() as db:
        saved = json.loads(
            db.execute("SELECT payload FROM reports WHERE id=?", (internal["id"],)).fetchone()[0]
        )
        assert saved["project"]["name"] == "Synthetic project"


def test_report_paginates_all_rows():
    p, times, rates = fixtures()
    snapshot = make_snapshot(
        {**p, "code": "SYN"},
        summarize(p, times, rates),
        [{"title": f"Item {i}", "kind": "issue", "priority": "high", "status": "open"} for i in range(21)],
        [],
        [],
    )
    prs = Presentation(io.BytesIO(make_pptx(snapshot, ["issues"])))
    texts = [
        cell.text
        for s in prs.slides
        for shape in s.shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    ]
    assert all(f"Item {i}" in texts for i in range(21))


def test_long_report_content_fits_and_is_retained():
    p, times, rates = fixtures()
    p.update(code="SYN", name="長專案名稱" * 25, summary="完整的專案摘要。" * 300)
    rows = [{"title": "議題內容" * 60, "kind": "issue", "priority": "high", "status": "open"}] * 5
    snapshot = make_snapshot(p, summarize(p, times, rates), rows, [], [], title="長報告標題" * 30)
    prs = Presentation(io.BytesIO(make_pptx(snapshot, ["issues"])))
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.top + shape.height <= prs.slide_height
    assert any(
        "完整的專案摘要" in shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert p["name"] in prs.slides[0].notes_slide.notes_text_frame.text
